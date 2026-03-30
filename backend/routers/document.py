from html import unescape
from io import BytesIO
from urllib.parse import urlparse
import re

import httpx
from fastapi import APIRouter, HTTPException, UploadFile, File
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document as DocxDocument


router = APIRouter(prefix="/api/documents", tags=["documents"])
MAX_PARSED_DOCUMENT_CHARS = 15000


class DocumentUrlRequest(BaseModel):
    url: str


def _is_valid_http_url(value: str) -> bool:
    try:
        parsed = urlparse(value)
        return parsed.scheme in {"http", "https"} and bool(parsed.netloc)
    except Exception:
        return False


def _extract_text_from_html(html: str) -> str:
    cleaned = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.IGNORECASE)
    cleaned = re.sub(r"<style[\s\S]*?</style>", " ", cleaned, flags=re.IGNORECASE)
    title_match = re.search(r"<title[^>]*>(.*?)</title>", cleaned, flags=re.IGNORECASE | re.DOTALL)
    title = unescape(title_match.group(1).strip()) if title_match else ""

    text = re.sub(r"<[^>]+>", " ", cleaned)
    text = unescape(text)
    text = re.sub(r"\s+", " ", text).strip()
    text = text[:12000]

    if title and title not in text:
      return f"页面标题：{title}\n\n页面正文：{text}"
    return f"页面正文：{text}"


def _has_meaningful_document_text(document: str) -> bool:
    if not document:
        return False

    stripped = document.strip()
    if not stripped:
        return False

    compact = re.sub(r"\s+", "", stripped)
    placeholders = {
        "页面正文：",
        "来源URL：",
    }
    if compact in placeholders:
        return False

    body_text = stripped.replace("页面标题：", "").replace("页面正文：", "").replace("来源 URL：", "").strip()
    return len(body_text) >= 20


def _extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    texts = []
    for page in reader.pages:
        texts.append(page.extract_text() or "")
    return "\n".join(texts).strip()[:MAX_PARSED_DOCUMENT_CHARS]


def _extract_text_from_docx(file_bytes: bytes) -> str:
    document = DocxDocument(BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return "\n".join(paragraphs).strip()[:MAX_PARSED_DOCUMENT_CHARS]


@router.post("/parse")
async def parse_uploaded_document(file: UploadFile = File(...)):
    file_name = file.filename or "未命名文档"
    lower_name = file_name.lower()
    content = await file.read()

    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空")

    if lower_name.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Excel 文档请在前端直接上传解析")

    try:
        if lower_name.endswith(".pdf"):
            extracted = _extract_text_from_pdf(content)
            doc_type = "PDF"
        elif lower_name.endswith(".docx"):
            extracted = _extract_text_from_docx(content)
            doc_type = "Word"
        elif lower_name.endswith(".doc"):
            raise HTTPException(status_code=400, detail="暂不支持 .doc，请先另存为 .docx")
        else:
            raise HTTPException(status_code=400, detail="仅支持 PDF、DOCX、XLSX、XLS 文件")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="文档解析失败，请检查文件内容") from exc

    full_document = f"文件名：{file_name}\n文档类型：{doc_type}\n\n{extracted}".strip()
    if not _has_meaningful_document_text(full_document):
        raise HTTPException(status_code=400, detail="未能从文档中提取到有效内容")

    return {
        "file_name": file_name,
        "document": full_document,
    }


@router.post("/fetch")
async def fetch_document_from_url(payload: DocumentUrlRequest):
    url = payload.url.strip()
    if not _is_valid_http_url(url):
        raise HTTPException(status_code=400, detail="请输入有效的 http/https URL")

    try:
        async with httpx.AsyncClient(follow_redirects=True, timeout=20.0) as client:
            response = await client.get(
                url,
                headers={
                    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                                  "AppleWebKit/537.36 (KHTML, like Gecko) "
                                  "Chrome/123.0.0.0 Safari/537.36"
                },
            )
            response.raise_for_status()
    except httpx.HTTPStatusError as exc:
        raise HTTPException(status_code=400, detail=f"文档 URL 抓取失败：HTTP {exc.response.status_code}") from exc
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=400, detail="文档 URL 抓取失败，请检查链接是否可访问") from exc

    content_type = response.headers.get("content-type", "").lower()
    body = response.text

    if "html" in content_type or "<html" in body.lower():
        document = _extract_text_from_html(body)
    else:
        document = body[:12000].strip()

    full_document = f"来源 URL：{url}\n\n{document}"
    if not _has_meaningful_document_text(full_document):
        raise HTTPException(
            status_code=400,
            detail="未能从该 URL 提取到有效文档内容，可能需要登录或没有访问权限"
        )

    return {
        "url": url,
        "title": urlparse(url).netloc,
        "document": full_document,
    }
